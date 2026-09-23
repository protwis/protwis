from collections import defaultdict

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

import openpyxl

from protein.models import (
    Gene,
    Protein,
    ProteinFamilyClassification,
    ProteinFamilyClassificationChemotype,
    ProteinFamilyClassificationModality,
    ProteinFamilyClassificationSense,
)


class Command(BaseCommand):
    help = "Import GPCR family classification annotations from Excel."

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._level4_cache = {}
        self._protein_resolution_cache = {}

    def add_arguments(self, parser):
        parser.add_argument(
            "-f",
            "--filename",
            action="store",
            dest="filename",
            help="Path to Classification.xlsx",
        )
        parser.add_argument(
            "-s",
            "--sheet",
            action="store",
            dest="sheet",
            default="Classification",
            help="Worksheet name to import",
        )

    def handle(self, *args, **options):
        filename = options.get("filename") or self.default_filename()
        sheet_name = options.get("sheet") or "Classification"

        data = self.parse_excel(filename, remove_header_linebreak=True)
        if sheet_name not in data:
            raise CommandError("Worksheet not found: {}".format(sheet_name))

        rows = data[sheet_name]
        self.stdout.write("Using file: {} (sheet: {})".format(filename, sheet_name))

        family_data, stats, seen_family_ids = self.aggregate_rows(rows)
        self.apply_updates(family_data, stats)
        self.prune_relationships(seen_family_ids, stats)

        self.stdout.write("Rows processed: {}".format(stats["rows_total"]))
        self.stdout.write("Gene matches: {}".format(stats["gene_matches"]))
        self.stdout.write("Entry short matches: {}".format(stats["entry_matches"]))
        self.stdout.write("Ambiguous skipped: {}".format(stats["ambiguous_skipped"]))
        self.stdout.write("Unmatched rows: {}".format(stats["unmatched"]))
        self.stdout.write("Families updated: {}".format(len(family_data)))
        self.stdout.write("Warnings: {}".format(stats["warnings"]))

    def default_filename(self):
        return settings.DATA_DIR + "/protein_data/Classification.xlsx"

    def aggregate_rows(self, rows):
        family_data = {}
        stats = defaultdict(int)
        seen_family_ids = set()

        for row in rows:
            stats["rows_total"] += 1
            gene_name = self.value_from_row(
                row, ["GPCRs (Gene name)", "GPCRs (Gene name)"]
            )
            uniprot_short = self.value_from_row(
                row, ["GPCRs (UniProt)", "GPCRs (UniProt)"]
            )

            protein, method, ambiguous, candidates = self.resolve_protein(
                gene_name, uniprot_short
            )
            if ambiguous:
                stats["warnings"] += 1
                stats["ambiguous_skipped"] += 1
                self.stdout.write(
                    "WARNING: ambiguous gene mapping for {} (candidates: {}).".format(
                        gene_name or "n/a", ", ".join(candidates) or "n/a"
                    )
                )
                continue
            if not protein:
                stats["unmatched"] += 1
                continue
            if method == "gene":
                stats["gene_matches"] += 1
            else:
                stats["entry_matches"] += 1

            family = self.resolve_level4_family(protein.family)
            if not family:
                stats["unmatched"] += 1
                continue
            seen_family_ids.add(family.id)

            if family.id not in family_data:
                family_data[family.id] = {
                    "family": family,
                    "senses": set(),
                    "chemotypes": {},
                    "modalities": {},
                    "order2_pairs": set(),
                    "row_keys": set(),
                }

            sense_value = self.value_from_row(row, ["Sense"])
            if sense_value:
                family_data[family.id]["senses"].add(sense_value)

            chemotype_value = self.value_from_row(row, ["Chemotype"])
            chemotype_order = self.value_from_row(row, ["Chemotype order"])
            chemotype_order_num = (
                self.parse_order(chemotype_order, family, "chemotype", stats)
                if chemotype_value
                else None
            )
            self.store_unique_order_value(
                family_data[family.id]["chemotypes"],
                chemotype_order_num,
                chemotype_value,
                family,
                "chemotype",
                stats,
            )

            modality_value = self.value_from_row(row, ["Modality"])
            modality_order = self.value_from_row(row, ["Modality order"])
            modality_order_num = (
                self.parse_order(modality_order, family, "modality", stats)
                if modality_value
                else None
            )
            self.store_unique_order_value(
                family_data[family.id]["modalities"],
                modality_order_num,
                modality_value,
                family,
                "modality",
                stats,
            )
            if chemotype_order_num == 2 or modality_order_num == 2:
                missing_fields = []
                if chemotype_order_num == 2 and not modality_value:
                    missing_fields.append("Modality")
                if modality_order_num == 2 and not chemotype_value:
                    missing_fields.append("Chemotype")
                if missing_fields:
                    stats["warnings"] += 1
                    self.stdout.write(
                        "WARNING: Incomplete Order 2 data for Receptor {}. Missing fields: {}. Falling back to Order 1 values.".format(
                            family.slug, ", ".join(missing_fields)
                        )
                    )
                family_data[family.id]["order2_pairs"].add(
                    (
                        chemotype_value if chemotype_order_num == 2 else None,
                        modality_value if modality_order_num == 2 else None,
                    )
                )

            row_key = (
                chemotype_value,
                chemotype_order_num,
                modality_value,
                modality_order_num,
                sense_value,
            )
            if row_key in family_data[family.id]["row_keys"]:
                stats["warnings"] += 1
                self.stdout.write(
                    "WARNING: duplicate row for {}: {}.".format(family.slug, row_key)
                )
            else:
                family_data[family.id]["row_keys"].add(row_key)

        return family_data, stats, seen_family_ids

    def apply_updates(self, family_data, stats):
        with transaction.atomic():
            for data in family_data.values():
                family = data["family"]

                # Sense logic: Non-sensory overrides all; Unknown is ignored if other values exist
                senses_to_apply = data["senses"]
                if "Non-sensory" in senses_to_apply:
                    senses_to_apply = {"Non-sensory"}
                elif "Unknown" in senses_to_apply and len(senses_to_apply) > 1:
                    senses_to_apply.remove("Unknown")

                sense_obj = None
                if senses_to_apply:
                    sense_name = sorted(senses_to_apply)[0]
                    sense_obj = self.get_or_create_vocab(
                        ProteinFamilyClassificationSense, sense_name
                    )

                chemotype_values = self.unique_values(
                    data["chemotypes"], data["order2_pairs"], index=0
                )
                default_chemotype = (
                    next(iter(chemotype_values)) if len(chemotype_values) == 1 else None
                )
                modality_values = self.unique_values(
                    data["modalities"], data["order2_pairs"], index=1
                )
                default_modality = (
                    next(iter(modality_values)) if len(modality_values) == 1 else None
                )

                # Replace rows for this family to keep grouped-by-order layout.
                ProteinFamilyClassification.objects.filter(
                    protein_family=family
                ).delete()

                rows_to_create = []
                order1_chemotype = data["chemotypes"].get(1) or default_chemotype
                order1_modality = data["modalities"].get(1) or default_modality

                if data["order2_pairs"] and not data["chemotypes"].get(1):
                    stats["warnings"] += 1
                    self.stdout.write(
                        "WARNING: order 2 chemotype without order 1 for {}.".format(
                            family.slug
                        )
                    )
                if data["order2_pairs"] and not data["modalities"].get(1):
                    stats["warnings"] += 1
                    self.stdout.write(
                        "WARNING: order 2 modality without order 1 for {}.".format(
                            family.slug
                        )
                    )
                if order1_chemotype or order1_modality:
                    rows_to_create.append((1, order1_chemotype, order1_modality))

                for chemotype_name, modality_name in sorted(data["order2_pairs"]):
                    chemotype_name = chemotype_name or default_chemotype
                    modality_name = modality_name or default_modality
                    if chemotype_name or modality_name:
                        rows_to_create.append((2, chemotype_name, modality_name))

                if rows_to_create:
                    for order_num, chemotype_name, modality_name in rows_to_create:
                        chemotype_obj = (
                            self.get_or_create_vocab(
                                ProteinFamilyClassificationChemotype, chemotype_name
                            )
                            if chemotype_name
                            else None
                        )
                        modality_obj = (
                            self.get_or_create_vocab(
                                ProteinFamilyClassificationModality, modality_name
                            )
                            if modality_name
                            else None
                        )
                        ProteinFamilyClassification.objects.get_or_create(
                            protein_family=family,
                            sense=sense_obj,
                            chemotype=chemotype_obj,
                            chemotype_order=order_num if chemotype_obj else None,
                            modality=modality_obj,
                            modality_order=order_num if modality_obj else None,
                        )
                elif sense_obj:
                    ProteinFamilyClassification.objects.get_or_create(
                        protein_family=family,
                        sense=sense_obj,
                    )

    def prune_relationships(self, seen_family_ids, stats):
        stale_rows = (
            ProteinFamilyClassification.objects.exclude(
                protein_family_id__in=seen_family_ids
            )
            if seen_family_ids
            else ProteinFamilyClassification.objects.all()
        )

        removed_count = stale_rows.count()

        if removed_count:
            stale_rows.delete()

        stats["pruned_total"] = removed_count

        self.stdout.write("Pruned stale annotations: total={}".format(removed_count))

    def parse_order(self, order_value, family, label, stats):
        try:
            order = int(order_value)
        except (TypeError, ValueError):
            stats["warnings"] += 1
            self.stdout.write(
                "WARNING: invalid {} order for {}.".format(label, family.slug)
            )
            return None
        if order not in (1, 2):
            stats["warnings"] += 1
            self.stdout.write(
                "WARNING: {} order out of range for {}.".format(label, family.slug)
            )
            return None
        return order

    def store_unique_order_value(
        self, storage, order, name_value, family, label, stats
    ):
        if not name_value or order != 1:
            return
        if order in storage and storage[order] != name_value:
            stats["warnings"] += 1
            self.stdout.write(
                "WARNING: conflicting {} order {} for {}. Existing: {}. New: {}.".format(
                    label, order, family.slug, storage[order], name_value
                )
            )
            return
        storage[order] = name_value

    def unique_values(self, order1_storage, order2_pairs, index):
        values = set()
        for value in order1_storage.values():
            if value:
                values.add(value)
        for pair in order2_pairs:
            value = pair[index]
            if value:
                values.add(value)
        return values

    def resolve_protein(self, gene_name, uniprot_short):
        cache_key = (gene_name or "", uniprot_short or "")
        if cache_key in self._protein_resolution_cache:
            return self._protein_resolution_cache[cache_key]

        if gene_name:
            genes = Gene.objects.filter(
                name__iexact=gene_name, species__common_name="Human"
            )
            proteins = (
                Protein.objects.filter(genes__in=genes, species__common_name="Human")
                .distinct()
                .order_by("id")
            )
            protein_count = proteins.count()
            if protein_count == 1:
                result = (proteins.first(), "gene", False, [])
                self._protein_resolution_cache[cache_key] = result
                return result
            if protein_count > 1:
                if uniprot_short:
                    short = uniprot_short.strip().lower()
                    short_matches = proteins.filter(
                        entry_name__istartswith="{}_".format(short)
                    )
                    if short_matches.count() == 1:
                        result = (short_matches.first(), "gene+entry", False, [])
                        self._protein_resolution_cache[cache_key] = result
                        return result
                primary_genes = genes.filter(position=0)
                primary_proteins = (
                    Protein.objects.filter(
                        genes__in=primary_genes, species__common_name="Human"
                    )
                    .distinct()
                    .order_by("id")
                )
                if primary_proteins.count() == 1:
                    result = (primary_proteins.first(), "gene+primary", False, [])
                    self._protein_resolution_cache[cache_key] = result
                    return result
                result = (
                    None,
                    "gene",
                    True,
                    list(proteins.values_list("entry_name", flat=True)),
                )
                self._protein_resolution_cache[cache_key] = result
                return result
            # No human match, fall back to any species
            genes_any = Gene.objects.filter(name__iexact=gene_name)
            proteins_any = (
                Protein.objects.filter(genes__in=genes_any).distinct().order_by("id")
            )
            any_count = proteins_any.count()
            if any_count == 1:
                result = (proteins_any.first(), "gene_any", False, [])
                self._protein_resolution_cache[cache_key] = result
                return result
            if any_count > 1:
                if uniprot_short:
                    short = uniprot_short.strip().lower()
                    short_matches = proteins_any.filter(
                        entry_name__istartswith="{}_".format(short)
                    )
                    if short_matches.count() == 1:
                        result = (short_matches.first(), "gene_any+entry", False, [])
                        self._protein_resolution_cache[cache_key] = result
                        return result
                primary_genes_any = genes_any.filter(position=0)
                primary_proteins_any = (
                    Protein.objects.filter(genes__in=primary_genes_any)
                    .distinct()
                    .order_by("id")
                )
                if primary_proteins_any.count() == 1:
                    result = (
                        primary_proteins_any.first(),
                        "gene_any+primary",
                        False,
                        [],
                    )
                    self._protein_resolution_cache[cache_key] = result
                    return result
                result = (
                    None,
                    "gene_any",
                    True,
                    list(proteins_any.values_list("entry_name", flat=True)),
                )
                self._protein_resolution_cache[cache_key] = result
                return result

        if uniprot_short:
            short = uniprot_short.strip().lower()
            proteins = Protein.objects.filter(
                entry_name__istartswith="{}_".format(short),
                species__common_name="Human",
            ).order_by("id")
            if proteins.exists():
                result = (proteins.first(), "entry", False, [])
                self._protein_resolution_cache[cache_key] = result
                return result

        result = (None, None, False, [])
        self._protein_resolution_cache[cache_key] = result
        return result

    def resolve_level4_family(self, family):
        if not family:
            return None
        if family.id in self._level4_cache:
            return self._level4_cache[family.id]

        lineage = []
        current = family
        while current:
            lineage.append(current)
            current = current.parent

        chain = list(reversed(lineage))
        if len(chain) < 5:
            resolved = chain[-1] if chain else None
        else:
            resolved = chain[4]

        self._level4_cache[family.id] = resolved
        return resolved

    def get_or_create_vocab(self, model, name):
        slug = slugify(name)
        obj, created = model.objects.get_or_create(slug=slug, defaults={"name": name})
        if not created and obj.name != name:
            obj.name = name
            obj.save(update_fields=["name"])
        return obj

    def value_from_row(self, row, keys):
        for key in keys:
            if key in row and row[key] != "":
                return row[key]
        return None

    def parse_excel(self, path, remove_header_linebreak=False):
        workbook = openpyxl.load_workbook(path, data_only=True)
        worksheets = workbook.sheetnames
        data = {}
        for worksheet_name in worksheets:
            if worksheet_name in data:
                continue

            data[worksheet_name] = []
            headers = []
            worksheet = workbook[worksheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            header_row = rows[0]
            num_cells = len(header_row)
            for col in range(num_cells):
                header = header_row[col]
                if not isinstance(header, str):
                    header = "" if header is None else str(header)
                if header == "":
                    header = "i_{}".format(col)
                if header in headers:
                    header += "_{}".format(col)
                if remove_header_linebreak:
                    header = header.replace("\n", " ")
                header = header.strip()
                headers.append(header)

            for row in rows[1:]:
                if not row:
                    continue
                key = row[0] if len(row) > 0 else None
                if key in ("", None):
                    continue
                row_dict = {}
                for col in range(num_cells):
                    row_dict[headers[col]] = row[col] if col < len(row) else None
                data[worksheet_name].append(row_dict)
        return data
